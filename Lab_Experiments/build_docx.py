#!/usr/bin/env python3
"""Build <folder name>_Report.docx and .pdf from REPORT.md in this folder.

The block between <!-- RUN START --> and <!-- RUN END --> in REPORT.md is replaced
with a fresh run of `python run_demo.py --markdown`, so the output shown in the
document is always the real output of the code.
House style: Times New Roman everywhere, black text, bordered tables, a title page,
1-inch margins and page numbers. Needs pandoc, python-docx and LibreOffice (soffice).
"""
import os
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.abspath(__file__))
MD = os.path.join(ROOT, "REPORT.md")
DOCX = os.path.join(ROOT, os.path.basename(ROOT) + "_Report.docx")
FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)


def force_font_and_colour(root):
    """Point every font at Times New Roman and every colour at black, overriding theme values."""
    for rf in root.iter(qn("w:rFonts")):
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rf.attrib.pop(qn(attr), None)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(attr), FONT)
    for c in root.iter(qn("w:color")):
        for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
            c.attrib.pop(qn(attr), None)
        c.set(qn("w:val"), "000000")


def set_style(doc, name, size, bold=None, italic=None, align=None, before=None, after=None):
    try:
        st = doc.styles[name]
    except KeyError:
        return
    st.font.name, st.font.size, st.font.color.rgb = FONT, Pt(size), BLACK
    if bold is not None:
        st.font.bold = bold
    if italic is not None:
        st.font.italic = italic
    if not hasattr(st, "paragraph_format"):  # character style
        return
    pf = st.paragraph_format
    if align is not None:
        pf.alignment = align
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)


def border_table(table):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        borders.append(e)
    tblPr.append(borders)
    width = tblPr.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        tblPr.append(width)
    width.set(qn("w:type"), "pct")
    width.set(qn("w:w"), "5000")
    for cell in table.rows[0].cells:  # header row in bold
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.size = Pt(11)


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.name, run.font.size, run.font.color.rgb = FONT, Pt(11), BLACK


def fill_run_output():
    start, end = "<!-- RUN START -->", "<!-- RUN END -->"
    with open(MD) as f:
        text = f.read()
    if start not in text:
        return
    run = subprocess.run([sys.executable, os.path.join(ROOT, "run_demo.py"), "--markdown"],
                         capture_output=True, text=True, check=True, cwd=ROOT).stdout
    head, rest = text.split(start, 1)
    with open(MD, "w") as f:
        f.write(head + start + "\n" + run + end + rest.split(end, 1)[1])


def main():
    fill_run_output()
    subprocess.run(["pandoc", MD, "-o", DOCX], check=True, cwd=ROOT)
    doc = Document(DOCX)

    set_style(doc, "Normal", 12)
    for body in ("Body Text", "First Paragraph"):
        set_style(doc, body, 12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=8)
    set_style(doc, "Compact", 11)
    set_style(doc, "Title", 20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=200, after=12)
    set_style(doc, "Subtitle", 15, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=60)
    set_style(doc, "Author", 13, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=4)
    set_style(doc, "Heading 1", 15, bold=True, italic=False, before=18, after=8)
    set_style(doc, "Heading 2", 13, bold=True, italic=False, before=14, after=6)
    set_style(doc, "Heading 3", 12, bold=True, italic=False, before=10, after=4)
    for name in ("Verbatim Char", "Hyperlink"):
        set_style(doc, name, 12)
    set_style(doc, "Source Code", 10, before=2, after=8)
    try:
        doc.styles["Source Code"].paragraph_format.left_indent = Inches(0.25)
    except KeyError:
        pass
    force_font_and_colour(doc.styles.element)
    force_font_and_colour(doc.element.body)

    for sec in doc.sections:
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
        add_page_number(sec)

    authors = [p for p in doc.paragraphs if p.style.name == "Author"]
    if authors:
        authors[0].paragraph_format.space_before = Pt(12)
        authors[-1].add_run().add_break(WD_BREAK.PAGE)

    for table in doc.tables:
        border_table(table)
    for tbl in doc.element.body.iter(qn("w:tbl")):
        nxt = tbl.getnext()
        if nxt is not None and nxt.tag == qn("w:p"):
            pPr = nxt.get_or_add_pPr()
            spacing = pPr.find(qn("w:spacing"))
            if spacing is None:
                spacing = OxmlElement("w:spacing")
                pPr.append(spacing)
            spacing.set(qn("w:before"), "200")
    doc.save(DOCX)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", ROOT, DOCX],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    print("built", os.path.basename(DOCX), "and PDF")


if __name__ == "__main__":
    main()
