#!/usr/bin/env python3
"""Build SmartCampus_MapReduce_Report.docx and .pdf from REPORT.md.

House style: Times New Roman everywhere, black text, bordered tables,
a separate title page, 1-inch margins and page numbers in the footer.
Needs pandoc, python-docx and LibreOffice (soffice).
"""
import os
import subprocess
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
MD = os.path.join(ROOT, "REPORT.md")
DOCX = os.path.join(ROOT, "SmartCampus_MapReduce_Report.docx")
FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)


def force_font_and_colour(root):
    """Point every font at Times New Roman and every colour at black, overriding theme values."""
    for rf in root.iter(qn("w:rFonts")):
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rf.attrib.pop(qn(attr), None)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(attr), FONT)
    for c in root.iter(qn("w:color")):
        for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
            c.attrib.pop(qn(attr), None)
        c.set(qn("w:val"), "000000")


def set_style(doc, name, size, bold=None, italic=None, align=None, before=None, after=None):
    try:
        st = doc.styles[name]
    except KeyError:
        return
    st.font.name, st.font.size, st.font.color.rgb = FONT, Pt(size), BLACK
    if bold is not None:
        st.font.bold = bold
    if italic is not None:
        st.font.italic = italic
    if not hasattr(st, "paragraph_format"):  # character style
        return
    pf = st.paragraph_format
    if align is not None:
        pf.alignment = align
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)


def border_table(table):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        borders.append(e)
    tblPr.append(borders)
    width = tblPr.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        tblPr.append(width)
    width.set(qn("w:type"), "pct")
    width.set(qn("w:w"), "5000")
    for cell in table.rows[0].cells:  # header row in bold
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.size = Pt(11)


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.name, run.font.size, run.font.color.rgb = FONT, Pt(11), BLACK


def fill_trace():
    """Replace the appendix between the TRACE markers with a fresh run of run_demo.py,
    followed by the verification and combiner checks, so the appendix always matches the code."""
    start, end = "<!-- TRACE START -->", "<!-- TRACE END -->"
    with open(MD) as f:
        text = f.read()
    if start not in text:
        return
    py = sys.executable
    trace = subprocess.run([py, os.path.join(ROOT, "run_demo.py"), "--markdown"],
                           capture_output=True, text=True, check=True, cwd=ROOT).stdout
    verify = subprocess.run([py, os.path.join(ROOT, "scripts", "verify.py")],
                            capture_output=True, text=True, check=True, cwd=ROOT).stdout
    comb = subprocess.run(["bash", os.path.join(ROOT, "scripts", "compare_combiner.sh")],
                          capture_output=True, text=True, check=True, cwd=ROOT).stdout
    extra = ("\n## Verification\n\n`scripts/verify.py` recomputes every answer with plain Python "
             "(no MapReduce) and compares it with the files in output/:\n\n```text\n"
             "$ python scripts/verify.py\n" + verify + "```\n\n"
             "## Combiner and Multiple Reducers\n\n`scripts/compare_combiner.sh` splits the Wi-Fi "
             "log into 4 map tasks, counts the records shuffled with and without the combiner, and "
             "runs Job 1 with 3 reducers:\n\n```text\n$ scripts/compare_combiner.sh\n" + comb + "```\n")
    head, rest = text.split(start, 1)
    tail = rest.split(end, 1)[1]
    with open(MD, "w") as f:
        f.write(head + start + "\n" + trace + extra + end + tail)


def main():
    fill_trace()
    subprocess.run(["pandoc", MD, "-o", DOCX], check=True)
    doc = Document(DOCX)

    set_style(doc, "Normal", 12)
    for body in ("Body Text", "First Paragraph"):
        set_style(doc, body, 12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=8)
    set_style(doc, "Compact", 11)
    set_style(doc, "Title", 20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=200, after=12)
    set_style(doc, "Subtitle", 15, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=60)
    set_style(doc, "Author", 13, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=4)
    set_style(doc, "Heading 1", 15, bold=True, italic=False, before=18, after=8)
    set_style(doc, "Heading 2", 13, bold=True, italic=False, before=14, after=6)
    set_style(doc, "Heading 3", 12, bold=True, italic=False, before=10, after=4)
    for name in ("Verbatim Char", "Hyperlink"):
        set_style(doc, name, 12)
    set_style(doc, "Source Code", 10, before=2, after=8)
    try:
        doc.styles["Source Code"].paragraph_format.left_indent = Inches(0.25)
    except KeyError:
        pass
    force_font_and_colour(doc.styles.element)
    force_font_and_colour(doc.element.body)

    for sec in doc.sections:
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
        add_page_number(sec)

    # Title page: break after the last author line (name, roll number, subgroup).
    authors = [p for p in doc.paragraphs if p.style.name == "Author"]
    if authors:
        authors[0].paragraph_format.space_before = Pt(12)
        authors[-1].add_run().add_break(WD_BREAK.PAGE)

    for table in doc.tables:
        border_table(table)
    # Breathing room between a table and the paragraph that follows it.
    for tbl in doc.element.body.iter(qn("w:tbl")):
        nxt = tbl.getnext()
        if nxt is not None and nxt.tag == qn("w:p"):
            pPr = nxt.get_or_add_pPr()
            spacing = pPr.find(qn("w:spacing"))
            if spacing is None:
                spacing = OxmlElement("w:spacing")
                pPr.append(spacing)
            spacing.set(qn("w:before"), "200")
    doc.save(DOCX)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", ROOT, DOCX],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    print("built", DOCX, "and PDF")


if __name__ == "__main__":
    main()
